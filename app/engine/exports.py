from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Any, Dict

from .advisor import generate_engineering_pack
from .models import ProjectIntake


def _pack(payload: Dict[str, Any] | ProjectIntake) -> Dict[str, Any]:
    return generate_engineering_pack(payload).model_dump()


def export_bom_csv(payload: Dict[str, Any] | ProjectIntake) -> str:
    pack = _pack(payload)
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow([
        "component_id", "categoria", "descripcion", "cantidad", "unidad", "proveedor",
        "marca", "modelo", "precio_unitario_usd", "precio_total_usd", "tipo_precio",
        "confianza", "semaforo", "priceguard_score", "accion_requerida", "stock", "entrega_dias", "vigencia", "nota_decision"
    ])
    req_map = {r["component_id"]: r for r in pack["requirements"]}
    for d in pack["market"]["price_decisions"]:
        r = req_map.get(d["component_id"], {})
        o = d.get("selected_offer") or {}
        writer.writerow([
            d["component_id"], r.get("category", ""), r.get("item", ""), d.get("qty", ""), r.get("unit", "u"),
            o.get("supplier_name", "RFQ requerido"), o.get("brand", ""), o.get("model", ""),
            d.get("unit_cost", 0), d.get("extended_cost", 0), d.get("price_type", ""),
            d.get("confidence_label", ""), d.get("semaphore_color", ""), d.get("priceguard_score", ""), d.get("action_required", ""), o.get("stock_status", "pendiente"), o.get("lead_time_days", ""),
            o.get("valid_until", ""), d.get("decision_note", "")
        ])
    return out.getvalue()


def export_bom_xlsx(payload: Dict[str, Any] | ProjectIntake) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    pack = _pack(payload)
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM cotizable"
    ws.append(["ControlPro Advisor OS V13 - BOM cotizable"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=17)
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="0B1722")
    ws.append([f"Proyecto: {pack['intake']['project_name']}"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=17)
    ws.append([])
    headers = ["ID", "Categoría", "Descripción", "Cant.", "Unidad", "Proveedor", "Marca", "Modelo", "Unit USD", "Total USD", "Tipo", "Confianza", "Semáforo", "PG Score", "Acción", "Stock", "Nota"]
    ws.append(headers)
    req_map = {r["component_id"]: r for r in pack["requirements"]}
    for d in pack["market"]["price_decisions"]:
        r = req_map.get(d["component_id"], {})
        o = d.get("selected_offer") or {}
        ws.append([
            d["component_id"], r.get("category", ""), r.get("item", ""), d.get("qty", ""), r.get("unit", "u"),
            o.get("supplier_name", "RFQ requerido"), o.get("brand", ""), o.get("model", ""),
            d.get("unit_cost", 0), d.get("extended_cost", 0), d.get("price_type", ""),
            d.get("confidence_label", ""), d.get("semaphore_color", ""), d.get("priceguard_score", ""), d.get("action_required", ""), o.get("stock_status", "pendiente"), d.get("decision_note", "")
        ])
    header_row = 4
    fill = PatternFill("solid", fgColor="123044")
    font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9E5EE")
    for cell in ws[header_row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    widths = [18,16,26,8,9,22,14,16,12,12,16,14,14,12,34,12,42]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A5"
    for row in ws.iter_rows(min_row=5, min_col=9, max_col=10):
        for cell in row:
            cell.number_format = '$#,##0.00'
    ws2 = wb.create_sheet("Resumen")
    summary = pack["market"]["summary"]
    budget = pack["budget"]
    rows = [
        ("Cobertura mercado", f"{summary['coverage_percent']}%"),
        ("Items con precio", f"{summary['items_with_price']}/{summary['total_items']}"),
        ("RFQ requeridos", summary["needs_rfq_count"]),
        ("PriceGuard", f"{summary.get('priceguard_score_percent', 0)}%"),
        ("Semáforo precios", f"Verde {summary.get('green_count',0)} / Amarillo {summary.get('yellow_count',0)} / Rojo {summary.get('red_count',0)}"),
        ("Costo materiales", budget["materials"]),
        ("Precio recomendado", budget["recommended_sell_price"]),
        ("Compuerta cotización", pack["quote_readiness"]["status"]),
        ("Compuerta construcción", pack["release_gates"]["construction_verdict"]),
    ]
    ws2.append(["Resumen ejecutivo", "Valor"])
    for row in rows:
        ws2.append(list(row))
    ws2["A1"].font = ws2["B1"].font = Font(bold=True, color="FFFFFF")
    ws2["A1"].fill = ws2["B1"].fill = PatternFill("solid", fgColor="0B1722")
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 58
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def export_pdf_report(payload: Dict[str, Any] | ProjectIntake) -> bytes:
    """Reporte PDF premium: portada, KPIs, PriceGuard, presupuesto, BOM y firmas.

    Es deliberadamente ejecutivo y trazable. No pretende ser plano CAD firmado.
    """
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

    pack = _pack(payload)
    bio = BytesIO()
    doc = SimpleDocTemplate(
        bio,
        pagesize=letter,
        rightMargin=0.52 * inch,
        leftMargin=0.52 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, leading=28, textColor=colors.HexColor("#06324A"), alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="Subtitle", parent=styles["BodyText"], fontSize=11, leading=15, textColor=colors.HexColor("#2B4B5B"), alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="H1Blue", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=19, textColor=colors.HexColor("#0B5C86"), spaceBefore=8, spaceAfter=6))
    styles.add(ParagraphStyle(name="H2Slate", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#123044"), spaceBefore=6, spaceAfter=4))
    styles.add(ParagraphStyle(name="BodySmall", parent=styles["BodyText"], fontSize=8.4, leading=10.2, textColor=colors.HexColor("#263B49")))
    styles.add(ParagraphStyle(name="Muted", parent=styles["BodyText"], fontSize=7.7, leading=9, textColor=colors.HexColor("#607380")))
    styles.add(ParagraphStyle(name="Right", parent=styles["BodyText"], alignment=TA_RIGHT, fontSize=8))

    def money(v: float) -> str:
        return f"${float(v):,.2f}"

    def para(txt: str, style="BodySmall"):
        return Paragraph(str(txt).replace("&", "&amp;"), styles[style])

    story = []
    story.append(Spacer(1, 0.25 * inch))
    story.append(Paragraph("ControlPro Advisor OS V13", styles["CoverTitle"]))
    story.append(Paragraph("FitLock Pro - Expediente técnico-comercial para cotización industrial", styles["Subtitle"]))
    story.append(Spacer(1, 0.22 * inch))
    cover_data = [
        ["Proyecto", pack["intake"]["project_name"]],
        ["Cliente", pack["intake"].get("client_name", "Cliente industrial")],
        ["Ubicación", f"{pack['intake']['location_city']}, {pack['intake']['location_province']}, {pack['intake']['country']}"],
        ["Aplicación", pack["intake"].get("application", "")],
        ["Motor", f"{pack['intake']['motor_power_hp']} HP · {pack['intake']['voltage']} V · {pack['intake']['phases']}F"],
        ["Versión / uso", "V13 FitLock Pro · Cotización y revisión profesional"],
    ]
    cover = Table(cover_data, colWidths=[1.55*inch, 5.05*inch])
    cover.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (0,-1), colors.HexColor("#0B1722")),
        ("TEXTCOLOR", (0,0), (0,-1), colors.white),
        ("BACKGROUND", (1,0), (1,-1), colors.HexColor("#F6FAFD")),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#C9D7E2")),
        ("FONTNAME", (0,0), (0,-1), "Helvetica-Bold"),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("PADDING", (0,0), (-1,-1), 8),
    ]))
    story.append(cover)
    story.append(Spacer(1, 0.18 * inch))
    story.append(para(pack["executive_verdict"]["verdict"]))
    story.append(Spacer(1, 0.12 * inch))
    notice = "Documento para cotización/revisión. No libera construcción ni energización sin placa real, campo, normativa aplicable, proveedor confirmado y aprobación humana."
    ntable = Table([[Paragraph("LÍMITE PROFESIONAL", styles["BodySmall"]), para(notice)]], colWidths=[1.55*inch, 5.05*inch])
    ntable.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (0,0), colors.HexColor("#8A5A00")),
        ("BACKGROUND", (1,0), (1,0), colors.HexColor("#FFF6DA")),
        ("TEXTCOLOR", (0,0), (0,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#E1BE62")),
        ("FONTNAME", (0,0), (0,0), "Helvetica-Bold"),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("PADDING", (0,0), (-1,-1), 8),
    ]))
    story.append(ntable)
    story.append(Spacer(1, 0.14 * inch))
    story.append(Paragraph("Índice ejecutivo", styles["H1Blue"]))
    index_rows = [
        ["1", "Resumen ejecutivo y KPIs"],
        ["2", "PriceGuard / confiabilidad de precios"],
        ["3", "Compuertas de calidad y riesgos"],
        ["4", "Solución recomendada y alternativas de arranque"],
        ["5", "Presupuesto comercial"],
        ["6", "BOM cotizable con semáforo"],
        ["7", "Supuestos, verificación y firmas"],
    ]
    it = Table(index_rows, colWidths=[0.45*inch, 5.95*inch])
    it.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D9E5EE")),
        ("BACKGROUND", (0,0), (0,-1), colors.HexColor("#EAF6FC")),
        ("FONTNAME", (0,0), (0,-1), "Helvetica-Bold"),
        ("PADDING", (0,0), (-1,-1), 6),
    ]))
    story.append(it)
    story.append(PageBreak())

    story.append(Paragraph("1. Resumen ejecutivo", styles["H1Blue"]))
    k = pack["statistics"]
    kpi_data = [
        ["Completitud ingeniería", "Cobertura catálogo", "PriceGuard", "Precio recomendado"],
        [f"{k['engineering_completeness_percent']}%", f"{k['market_coverage_percent']}%", f"{k.get('priceguard_score_percent', 0)}%", money(pack['budget']['recommended_sell_price'])],
    ]
    kpi = Table(kpi_data, colWidths=[1.65*inch]*4)
    kpi.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0B1722")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("BACKGROUND", (0,1), (-1,1), colors.HexColor("#EAF6FC")),
        ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#CAD7E0")),
        ("ALIGN", (0,0), (-1,-1), "CENTER"),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTNAME", (0,1), (-1,1), "Helvetica-Bold"),
        ("FONTSIZE", (0,1), (-1,1), 13),
        ("PADDING", (0,0), (-1,-1), 7),
    ]))
    story.append(kpi)
    story.append(Spacer(1, 9))
    story.append(Paragraph("2. Semáforo PriceGuard", styles["H1Blue"]))
    ms = pack["market"]["summary"]
    pg_rows = [
        ["Métrica", "Valor", "Lectura"],
        ["Verde", ms.get("green_count", 0), para("Precio usable para cotización revisable.")],
        ["Amarillo", ms.get("yellow_count", 0), para("Referencial; conviene RFQ/confirmación.")],
        ["Rojo", ms.get("red_count", 0), para("No cerrar precio; exige RFQ/autocorrección.")],
        ["Autocorrecciones", ms.get("auto_corrected_count", 0), para("La app detectó valores fuera de banda o débiles.")],
        ["Regla de verdad", "-", para(ms.get("truth_status", "Precio final solo con proveedor confirmado."))],
    ]
    pg_table = Table(pg_rows, colWidths=[1.55*inch, 1.2*inch, 3.95*inch])
    pg_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#123044")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#D9E5EE")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("BACKGROUND", (0,1), (-1,1), colors.HexColor("#E7FFF2")),
        ("BACKGROUND", (0,2), (-1,2), colors.HexColor("#FFF6DA")),
        ("BACKGROUND", (0,3), (-1,3), colors.HexColor("#FFECEC")),
    ]))
    story.append(pg_table)
    story.append(Spacer(1, 7))
    audit = ms.get("candidate_offer_audit", {})
    audit_rows = [
        ["Ofertas auditadas", audit.get("total_candidate_offers", 0)],
        ["Bajos rechazados", audit.get("suspicious_low_rejected", 0)],
        ["Altos marcados", audit.get("suspicious_high_flagged", 0)],
        ["Fuentes débiles", audit.get("weak_source_flagged", 0)],
        ["Veredicto", ms.get("priceguard_verdict", "Revisable")],
    ]
    story.append(Paragraph("2.1 Auditoría anti-precios basura", styles["H2Slate"]))
    audit_t = Table(audit_rows, colWidths=[2.2*inch, 4.2*inch])
    audit_t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (0,-1), colors.HexColor("#0B1722")),
        ("TEXTCOLOR", (0,0), (0,-1), colors.white),
        ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#D9E5EE")),
        ("PADDING", (0,0), (-1,-1), 6),
    ]))
    story.append(audit_t)
    story.append(Spacer(1, 7))
    story.append(para(ms.get("truth_status", "Precio cerrado solo con proveedor confirmado.")))
    story.append(para(ms.get("catalog_scope", "Catálogo piloto y RFQ.")))

    story.append(Paragraph("3. Compuertas de calidad", styles["H1Blue"]))
    gate_rows = [["Compuerta", "Estado", "Acción requerida"]]
    for g in pack["release_gates"]["gates"]:
        gate_rows.append([para(g["name"]), "OK" if g["passed"] else "PENDIENTE", para(g["required_action"])])
    gate_t = Table(gate_rows, colWidths=[1.7*inch, 1.0*inch, 4.0*inch])
    gate_t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#123044")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#D9E5EE")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
    ]))
    story.append(gate_t)

    story.append(Paragraph("4. Solución recomendada", styles["H1Blue"]))
    rec = pack["recommended_option"]
    story.append(para(f"{rec['name']} - {rec.get('fit','')}. {rec.get('why','')}"))
    story.append(para(f"Cómo funciona: {rec.get('how_it_works','')}"))
    story.append(para(f"Impacto en precio/riesgo: {rec.get('price_impact','')}"))
    alt_rows = [["Alternativa", "Costo", "Control", "Seguridad", "Cuándo conviene"]]
    for o in pack.get("alternatives", [])[:6]:
        alt_rows.append([para(o.get("name", "")), o.get("initial_cost", ""), f"{o.get('control_quality',0)}%", f"{o.get('safety_depth',0)}%", para(o.get("better_when", o.get("sell_when", "")))])
    alt_t = Table(alt_rows, colWidths=[1.55*inch, 0.7*inch, 0.65*inch, 0.75*inch, 3.0*inch], repeatRows=1)
    alt_t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#123044")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D9E5EE")),
        ("FONTSIZE", (0,0), (-1,-1), 6.8),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
    ]))
    story.append(Spacer(1, 8))
    story.append(alt_t)

    story.append(Paragraph("5. Presupuesto", styles["H1Blue"]))
    budget_names = [
        ("Materiales", "materials"), ("Armado tablero", "panel_labor"), ("Instalación campo", "field_labor"),
        ("Ingeniería", "engineering"), ("Transporte/logística", "transport_logistics"), ("Contingencia", "contingency"),
        ("Margen", "margin"), ("Precio piso", "floor_price"), ("Precio recomendado", "recommended_sell_price"), ("Precio premium", "premium_price"),
    ]
    budget_rows = [["Concepto", "Valor USD"]] + [[name, money(pack["budget"].get(key, 0))] for name, key in budget_names]
    bt = Table(budget_rows, colWidths=[3.4*inch, 2.3*inch])
    bt.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#123044")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#D9E5EE")),
        ("ALIGN", (1,1), (1,-1), "RIGHT"),
        ("FONTNAME", (0,-2), (-1,-1), "Helvetica-Bold"),
        ("BACKGROUND", (0,-2), (-1,-2), colors.HexColor("#EAF6FC")),
        ("BACKGROUND", (0,-1), (-1,-1), colors.HexColor("#E7FFF2")),
    ]))
    story.append(bt)
    story.append(PageBreak())

    story.append(Paragraph("6. BOM cotizable con semáforo", styles["H1Blue"]))
    req_map = {r["component_id"]: r for r in pack["requirements"]}
    rows = [["Componente", "Proveedor / Modelo", "Semáforo", "Unit", "Total", "Acción"]]
    for d in pack["market"]["price_decisions"][:22]:
        offer = d.get("selected_offer") or {}
        r = req_map.get(d["component_id"], {})
        rows.append([
            para(f"{r.get('item', d['component_id'])}<br/>{d['component_id']}"),
            para(f"{offer.get('supplier_name','RFQ')}<br/>{offer.get('brand','')} {offer.get('model','')}"),
            para(f"{d.get('semaphore_color','')}<br/>PG {d.get('priceguard_score',0)}%"),
            money(d.get("unit_cost", 0)),
            money(d.get("extended_cost", 0)),
            para(d.get("action_required", "")),
        ])
    bom = Table(rows, colWidths=[1.55*inch, 1.6*inch, 0.9*inch, 0.78*inch, 0.78*inch, 1.42*inch], repeatRows=1)
    bom.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#123044")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D9E5EE")),
        ("FONTSIZE", (0,0), (-1,-1), 6.7),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ALIGN", (3,1), (4,-1), "RIGHT"),
    ]))
    story.append(bom)
    story.append(Spacer(1, 8))
    story.append(Paragraph("7. Supuestos y verificación", styles["H1Blue"]))
    assumption_rows = [["Tema", "Supuesto", "Confianza", "Cómo verificar"]]
    for a in pack.get("assumption_ledger", [])[:10]:
        assumption_rows.append([para(a["topic"]), para(a["assumption"]), a["confidence"], para(a["verification"])])
    at = Table(assumption_rows, colWidths=[1.25*inch, 2.3*inch, 0.9*inch, 2.25*inch], repeatRows=1)
    at.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#123044")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D9E5EE")),
        ("FONTSIZE", (0,0), (-1,-1), 6.8),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
    ]))
    story.append(at)
    story.append(Spacer(1, 8))
    story.append(Paragraph("8. Firmas y control de revisión", styles["H1Blue"]))
    sign = Table([
        ["Elaborado por", "Revisado por", "Aprobación cliente"],
        ["\n\n________________________\nNombre / fecha", "\n\n________________________\nResponsable técnico / fecha", "\n\n________________________\nNombre / fecha"],
    ], colWidths=[2.15*inch, 2.15*inch, 2.15*inch])
    sign.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0B1722")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#C9D7E2")),
        ("ALIGN", (0,0), (-1,-1), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
    ]))
    story.append(sign)
    story.append(Spacer(1, 8))
    story.append(para(pack["human_review_notice"], "Muted"))

    def _footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#607380"))
        canvas.drawString(0.55*inch, 0.35*inch, "ControlPro Advisor OS V13 FitLock Pro - Documento para cotización/revisión")
        canvas.drawRightString(7.95*inch, 0.35*inch, f"Página {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return bio.getvalue()
